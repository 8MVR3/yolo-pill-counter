"""Шаг 2. Считает объекты на всех фото в папке.

Сохраняет копии фото с рамками и подписью «Найдено: N» и таблицу Excel
с количеством объектов по каждому фото.

Запуск (без папки возьмёт проверочные фото датасета):
    python count.py runs/pills/weights/best.pt
    python count.py runs/pills/weights/best.pt папка_с_фото
"""
import sys
from collections import Counter
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from ultralytics import YOLO

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}
OUT = Path("results")


def load_font(size):
    for name in ("arial.ttf", "C:/Windows/Fonts/arial.ttf", "DejaVuSans.ttf",
                 "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def default_source():
    from ultralytics.data.utils import check_det_dataset
    return Path(check_det_dataset("medical-pills.yaml")["val"])


def main(weights, source=None):
    source = Path(source) if source else default_source()
    images = sorted(p for p in source.iterdir() if p.suffix.lower() in IMAGE_EXT)
    model = YOLO(weights)
    names = [model.names[i] for i in sorted(model.names)]
    OUT.mkdir(exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Подсчёт"
    ws.append(["Фото", *names, "Всего"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="5B3CC4")
    ws.column_dimensions["A"].width = 32

    for path in images:
        res = model.predict(str(path), conf=0.4, imgsz=640, verbose=False)[0]
        counts = Counter(model.names[int(c)] for c in res.boxes.cls.tolist())
        total = sum(counts.values())

        plotted = Image.fromarray(np.ascontiguousarray(res.plot(labels=False, conf=False, line_width=2)[..., ::-1]))
        d = ImageDraw.Draw(plotted)
        font = load_font(max(18, plotted.width // 28))
        text = f"Найдено: {total}"
        box = d.textbbox((0, 0), text, font=font)
        pad = font.size // 2
        d.rectangle((0, 0, box[2] + 2 * pad, box[3] + 2 * pad), fill=(91, 60, 196))
        d.text((pad, pad), text, font=font, fill="white")
        plotted.save(OUT / path.name)

        ws.append([path.name, *[counts.get(n, 0) for n in names], total])
        print(f"{path.name}: {total}")

    last = ws.max_row
    ws.append(["Итого", *[f"=SUM({chr(66 + i)}2:{chr(66 + i)}{last})" for i in range(len(names) + 1)]])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    wb.save(OUT / "Подсчёт.xlsx")
    print(f"\nГотово: фото с рамками и таблица в папке {OUT.resolve()}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    main(*sys.argv[1:3])
